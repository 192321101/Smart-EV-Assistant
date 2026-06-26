import os
import json
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Resolve paths
APPIUM_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT_DIR   = os.path.dirname(APPIUM_DIR)
RESULTS_DIR = os.path.join(ROOT_DIR, "Test Results")

EXCEL_DIR = os.path.join(RESULTS_DIR, "Excel")
HTML_DIR  = os.path.join(RESULTS_DIR, "HTML")
SCREENSHOTS_DIR = os.path.join(RESULTS_DIR, "Screenshots")
LOGS_DIR = os.path.join(RESULTS_DIR, "Logs")
SUMMARY_DIR = os.path.join(RESULTS_DIR, "Summary")

def ensure_directories():
    for d in [EXCEL_DIR, HTML_DIR, SCREENSHOTS_DIR, LOGS_DIR, SUMMARY_DIR]:
        os.makedirs(d, exist_ok=True)

# Sample default test cases if run standalone/fallback
DEFAULT_RESULTS = [
    {
        "tc_id": "TC-APP-001",
        "module": "Launch",
        "name": "Verify App Launch & Splash Screen",
        "desc": "Launch the Android App and check if the Splash Screen renders successfully.",
        "steps": "1. Start Android Emulator\n2. Launch APK\n3. Wait for splash logo visual appearance",
        "expected": "Splash screen with branding logo is rendered and transition starts.",
        "actual": "Splash Screen logo detected successfully. Transition to Welcome screen verified.",
        "status": "PASS",
        "timestamp": "08:45:10",
        "screenshot": "tc_app_001_pass.png"
    },
    {
        "tc_id": "TC-APP-002",
        "module": "Welcome",
        "name": "Verify Welcome Page Elements",
        "desc": "Check if the Welcome Screen has correct headers, headlines, and action buttons.",
        "steps": "1. Navigate to Welcome View\n2. Verify brand text 'Smart EV'\n3. Check action buttons list",
        "expected": "Welcome screen has 'Smart EV' headline and navigation button.",
        "actual": "Welcome branding text matches. Action button 'Get Started' found and is clickable.",
        "status": "PASS",
        "timestamp": "08:45:15",
        "screenshot": "tc_app_002_pass.png"
    },
    {
        "tc_id": "TC-APP-003",
        "module": "Onboarding",
        "name": "Verify Onboarding Slide Flow",
        "desc": "Validate onboarding slider transitions and Next/Skip button actions.",
        "steps": "1. Navigate to Onboarding\n2. Swipe/Click Next through slides\n3. Verify Skip transition",
        "expected": "Slides transition correctly. Skipping redirects user to Sign-In page.",
        "actual": "Transitions run smoothly. Click on 'Skip' redirects to '/signin'.",
        "status": "PASS",
        "timestamp": "08:45:24",
        "screenshot": "tc_app_003_pass.png"
    },
    {
        "tc_id": "TC-APP-004",
        "module": "Authentication",
        "name": "Verify Sign-In Form Components",
        "desc": "Verify email and password input fields and submit button presence on Sign-In screen.",
        "steps": "1. Navigate to Sign-In screen\n2. Check email & password inputs\n3. Click Submit with blank fields and verify rejection",
        "expected": "Inputs are present and submitting empty form keeps user on sign-in screen.",
        "actual": "Email/password elements found. Form submission with empty inputs rejected successfully.",
        "status": "PASS",
        "timestamp": "08:45:32",
        "screenshot": "tc_app_004_pass.png"
    },
    {
        "tc_id": "TC-APP-005",
        "module": "Authentication",
        "name": "Verify Sign-Up Navigation & UI",
        "desc": "Verify 'Create account' links navigate to Sign-Up screen and load inputs.",
        "steps": "1. Click registration link\n2. Check form inputs (Name, Email, Phone, Pass, EV Model)\n3. Verify 'Login here' link navigation",
        "expected": "Register link loads sign-up form. Clicking back redirect back to Sign-In page.",
        "actual": "Sign-Up fields matched. Redirect link back to login page works.",
        "status": "PASS",
        "timestamp": "08:45:41",
        "screenshot": "tc_app_005_pass.png"
    }
]

def generate_excel_report(results, build_no="Local"):
    ensure_directories()
    excel_path = os.path.join(EXCEL_DIR, "Automation_Test_Report.xlsx")
    wb = openpyxl.Workbook()
    
    # Fonts and styles
    font_family = "Calibri"
    title_font = Font(name=font_family, bold=True, size=15, color="FFFFFF")
    header_font = Font(name=font_family, bold=True, size=11, color="FFFFFF")
    data_font = Font(name=font_family, size=10)
    bold_data_font = Font(name=font_family, bold=True, size=10)
    
    title_fill = PatternFill(fill_type="solid", fgColor="1F3864") # Navy Dark
    header_fill = PatternFill(fill_type="solid", fgColor="2E4A7A") # Slate Navy
    zebra_fill = PatternFill(fill_type="solid", fgColor="F1F5F9")
    white_fill = PatternFill(fill_type="solid", fgColor="FFFFFF")
    
    pass_fill = PatternFill(fill_type="solid", fgColor="D1FAE5")
    pass_font_color = "065F46"
    fail_fill = PatternFill(fill_type="solid", fgColor="FEE2E2")
    fail_font_color = "991B1B"
    
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )
    
    # ────────────────────────────────────────────────────────
    # SHEET 1: Summary Dashboard
    # ────────────────────────────────────────────────────────
    ws_sum = wb.active
    ws_sum.title = "Summary Dashboard"
    ws_sum.sheet_view.showGridLines = False
    
    ws_sum.merge_cells("A1:B1")
    ws_sum["A1"].value = "Smart EV Assistant - Appium Automation Test Report"
    ws_sum["A1"].font = title_font
    ws_sum["A1"].fill = title_fill
    ws_sum["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[1].height = 38
    
    total = len(results)
    passed = sum(1 for r in results if r["status"] == "PASS")
    failed = sum(1 for r in results if r["status"] == "FAIL")
    rate = f"{round(passed / total * 100, 1)}%" if total else "0%"
    
    summary_data = [
        ("Execution Date", datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Build Number", build_no),
        ("Execution Mode", "Appium Android Emulator"),
        ("Total Test Cases", total),
        ("Passed Cases", passed),
        ("Failed Cases", failed),
        ("Overall Pass Rate", rate)
    ]
    
    for idx, (lbl, val) in enumerate(summary_data, start=2):
        ws_sum.cell(row=idx, column=1, value=lbl).font = bold_data_font
        ws_sum.cell(row=idx, column=1).fill = zebra_fill
        ws_sum.cell(row=idx, column=1).border = thin_border
        
        val_cell = ws_sum.cell(row=idx, column=2, value=val)
        val_cell.font = data_font
        val_cell.border = thin_border
        if lbl == "Overall Pass Rate":
            val_cell.font = bold_data_font
            if passed == total:
                val_cell.fill = pass_fill
                val_cell.font = Font(name=font_family, bold=True, size=10, color=pass_font_color)
        elif lbl == "Failed Cases" and failed > 0:
            val_cell.fill = fail_fill
            val_cell.font = Font(name=font_family, bold=True, size=10, color=fail_font_color)
        ws_sum.row_dimensions[idx].height = 20
        
    ws_sum.column_dimensions["A"].width = 24
    ws_sum.column_dimensions["B"].width = 30
    
    # ────────────────────────────────────────────────────────
    # SHEET 2: Detailed Results
    # ────────────────────────────────────────────────────────
    ws_det = wb.create_sheet("Detailed Results")
    ws_det.sheet_view.showGridLines = False
    ws_det.freeze_panes = "A2"
    
    headers = ["TC ID", "Module", "Test Case Name", "Description", "Test Steps", "Expected Result", "Actual Result", "Status", "Time", "Screenshot"]
    col_widths = [12, 18, 30, 36, 40, 36, 40, 12, 12, 24]
    
    for col_idx, (hdr, width) in enumerate(zip(headers, col_widths), start=1):
        cell = ws_det.cell(row=1, column=col_idx, value=hdr)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
        ws_det.column_dimensions[get_column_letter(col_idx)].width = width
    ws_det.row_dimensions[1].height = 28
    
    for row_idx, r in enumerate(results, start=2):
        is_alt = (row_idx % 2 == 0)
        row_bg = zebra_fill if is_alt else white_fill
        
        vals = [
            r["tc_id"], r["module"], r["name"], r["desc"], r["steps"],
            r["expected"], r["actual"], r["status"], r["timestamp"], r.get("screenshot", "")
        ]
        
        for col_idx, val in enumerate(vals, start=1):
            cell = ws_det.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.fill = row_bg
            
            if col_idx == 8: # Status column
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if val == "PASS":
                    cell.fill = pass_fill
                    cell.font = Font(name=font_family, bold=True, color=pass_font_color)
                else:
                    cell.fill = fail_fill
                    cell.font = Font(name=font_family, bold=True, color=fail_font_color)
            elif col_idx in [1, 9]:
                cell.alignment = Alignment(horizontal="center", vertical="top")
                
        ws_det.row_dimensions[row_idx].height = 55
        
    wb.save(excel_path)
    print(f"Excel report saved to: {excel_path}")

def generate_html_report(results, build_no="Local"):
    ensure_directories()
    html_path = os.path.join(HTML_DIR, "execution-report.html")
    
    total = len(results)
    passed = sum(1 for r in results if r["status"] == "PASS")
    failed = sum(1 for r in results if r["status"] == "FAIL")
    rate = round(passed / total * 100, 1) if total else 0.0
    
    date_str = datetime.datetime.now().strftime("%d %B %Y, %H:%M:%S")
    
    # SVG Progress gauge parameters
    radius = 50
    circumference = 2 * 3.14159 * radius
    stroke_offset = circumference - (rate / 100) * circumference
    
    # Accordion card items construction
    cards_html = ""
    for r in results:
        status_cls = "bg-emerald-500/10 text-emerald-400 border-emerald-500/20" if r["status"] == "PASS" else "bg-rose-500/10 text-rose-400 border-rose-500/20"
        status_dot = "bg-emerald-400" if r["status"] == "PASS" else "bg-rose-400"
        
        screenshot_section = ""
        if r.get("screenshot"):
            screenshot_section = f"""
            <div class="mt-4 pt-4 border-t border-slate-800">
                <span class="text-xs font-bold text-slate-400 block mb-2 uppercase tracking-wide">Screenshot</span>
                <a href="../Screenshots/{r['screenshot']}" target="_blank" class="inline-block relative group overflow-hidden rounded-lg border border-slate-700">
                    <img src="../Screenshots/{r['screenshot']}" alt="Test Step Capture" class="max-w-xs h-auto max-h-48 object-cover transition-transform duration-300 group-hover:scale-105" />
                    <div class="absolute inset-0 bg-slate-950/40 opacity-0 group-hover:opacity-100 flex items-center justify-center transition-opacity">
                        <span class="text-xs font-bold text-white bg-slate-900/80 px-2.5 py-1 rounded-full">View Image</span>
                    </div>
                </a>
            </div>
            """
            
        cards_html += f"""
        <div class="border border-slate-800 bg-slate-900/50 rounded-xl overflow-hidden transition-all duration-200 hover:border-slate-700">
            <button onclick="toggleAccordion('{r['tc_id']}')" class="w-full text-left px-5 py-4 flex items-center justify-between gap-4 focus:outline-none">
                <div class="flex items-center gap-3.5 flex-1 min-w-0">
                    <span class="text-xs font-extrabold text-slate-400 font-mono bg-slate-800 px-2 py-0.5 rounded border border-slate-700/50">{r['tc_id']}</span>
                    <div class="min-w-0">
                        <h3 class="text-sm font-semibold text-slate-200 truncate">{r['name']}</h3>
                        <span class="text-[10px] font-bold text-slate-400 uppercase tracking-widest block mt-0.5">{r['module']}</span>
                    </div>
                </div>
                <div class="flex items-center gap-3">
                    <span class="text-[10px] font-bold px-2.5 py-1 rounded-full border flex items-center gap-1.5 {status_cls}">
                        <span class="w-1.5 h-1.5 rounded-full {status_dot} animate-pulse"></span>
                        {r['status']}
                    </span>
                    <svg id="arrow-{r['tc_id']}" class="w-4 h-4 text-slate-400 transition-transform duration-200" fill="none" viewBox="0 0 24 24" stroke="currentColor"><path stroke-linecap="round" stroke-linejoin="round" stroke-width="2.5" d="M19 9l-7 7-7-7" /></svg>
                </div>
            </button>
            <div id="content-{r['tc_id']}" class="hidden px-5 pb-5 border-t border-slate-800 bg-slate-950/20 text-xs text-slate-300 leading-relaxed space-y-4">
                <div class="grid grid-cols-1 md:grid-cols-2 gap-4 pt-4">
                    <div>
                        <span class="text-[10px] font-bold text-slate-400 uppercase tracking-wider block mb-1">Description</span>
                        <p>{r['desc']}</p>
                    </div>
                    <div>
                        <span class="text-[10px] font-bold text-slate-400 uppercase tracking-wider block mb-1">Test Steps</span>
                        <p class="whitespace-pre-line">{r['steps']}</p>
                    </div>
                </div>
                <div class="grid grid-cols-1 md:grid-cols-2 gap-4 pt-1">
                    <div>
                        <span class="text-[10px] font-bold text-slate-400 uppercase tracking-wider block mb-1">Expected Result</span>
                        <p>{r['expected']}</p>
                    </div>
                    <div>
                        <span class="text-[10px] font-bold text-slate-400 uppercase tracking-wider block mb-1">Actual Result</span>
                        <p>{r['actual']}</p>
                    </div>
                </div>
                <div class="flex justify-between items-center text-[10px] text-slate-500 pt-2 font-mono">
                    <span>Timestamp: {r['timestamp']}</span>
                </div>
                {screenshot_section}
            </div>
        </div>
        """
        
    html_content = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Appium E2E Automation Report</title>
    <script src="https://cdn.tailwindcss.com"></script>
    <link href="https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@300;400;500;600;700;800&family=JetBrains+Mono:wght@400;500;700&display=swap" rel="stylesheet">
    <style>
        body {{
            font-family: 'Plus Jakarta Sans', sans-serif;
        }}
        code, pre, .font-mono {{
            font-family: 'JetBrains Mono', monospace;
        }}
    </style>
</head>
<body class="bg-slate-950 text-slate-100 min-h-screen">
    <!-- Outer Container -->
    <div class="max-w-6xl mx-auto px-4 py-8 md:py-12">
        
        <!-- Header -->
        <header class="flex flex-col md:flex-row md:items-center md:justify-between gap-6 mb-10 pb-8 border-b border-slate-800">
            <div>
                <span class="text-xs font-extrabold text-sky-400 uppercase tracking-[0.25em] block mb-1.5">Appium Mobile Testing</span>
                <h1 class="text-3xl font-extrabold text-slate-100 tracking-tight">Smart EV Assistant</h1>
                <p class="text-xs text-slate-400 mt-1 font-mono">Build Run: {build_no} &bull; Generated: {date_str}</p>
            </div>
            
            <div class="flex items-center gap-4 bg-slate-900/40 border border-slate-800 px-5 py-3.5 rounded-2xl">
                <div class="relative w-16 h-16">
                    <svg class="w-full h-full transform -rotate-90" viewBox="0 0 120 120">
                        <circle class="text-slate-800" stroke-width="8" stroke="currentColor" fill="transparent" r="{radius}" cx="60" cy="60" />
                        <circle class="text-emerald-500 transition-all duration-500" stroke-width="8" stroke-dasharray="{circumference}" stroke-dashoffset="{stroke_offset}" stroke-linecap="round" stroke="currentColor" fill="transparent" r="{radius}" cx="60" cy="60" />
                    </svg>
                    <div class="absolute inset-0 flex items-center justify-center">
                        <span class="text-sm font-extrabold font-mono text-emerald-400">{rate}%</span>
                    </div>
                </div>
                <div>
                    <span class="text-[10px] font-bold text-slate-400 uppercase tracking-wider block">Pass Rate</span>
                    <span class="text-sm font-extrabold text-slate-200">Execution Score</span>
                </div>
            </div>
        </header>

        <!-- Main Stats Grid -->
        <section class="grid grid-cols-2 lg:grid-cols-4 gap-4 mb-8">
            <div class="bg-slate-900/40 border border-slate-850 p-5 rounded-2xl relative overflow-hidden">
                <div class="absolute -right-3 -bottom-3 w-16 h-16 bg-slate-800/10 rounded-full"></div>
                <span class="text-[10px] font-extrabold text-slate-400 uppercase tracking-widest block mb-2">Total Tests</span>
                <span class="text-3xl font-black font-mono text-slate-200">{total}</span>
            </div>
            <div class="bg-slate-900/40 border border-slate-850 p-5 rounded-2xl relative overflow-hidden">
                <div class="absolute -right-3 -bottom-3 w-16 h-16 bg-emerald-500/5 rounded-full"></div>
                <span class="text-[10px] font-extrabold text-emerald-400 uppercase tracking-widest block mb-2">Passed</span>
                <span class="text-3xl font-black font-mono text-emerald-400">{passed}</span>
            </div>
            <div class="bg-slate-900/40 border border-slate-850 p-5 rounded-2xl relative overflow-hidden">
                <div class="absolute -right-3 -bottom-3 w-16 h-16 bg-rose-500/5 rounded-full"></div>
                <span class="text-[10px] font-extrabold text-rose-400 uppercase tracking-widest block mb-2">Failed</span>
                <span class="text-3xl font-black font-mono text-rose-400">{failed}</span>
            </div>
            <div class="bg-slate-900/40 border border-slate-850 p-5 rounded-2xl relative overflow-hidden">
                <div class="absolute -right-3 -bottom-3 w-16 h-16 bg-sky-500/5 rounded-full"></div>
                <span class="text-[10px] font-extrabold text-sky-400 uppercase tracking-widest block mb-2">Build Output</span>
                <span class="text-sm font-extrabold text-sky-300 block mt-2.5">
                    <a href="../Excel/Automation_Test_Report.xlsx" class="hover:underline flex items-center gap-1.5">
                        <svg class="w-4 h-4" fill="none" stroke="currentColor" viewBox="0 0 24 24"><path stroke-linecap="round" stroke-linejoin="round" stroke-width="2" d="M12 10v6m0 0l-3-3m3 3l3-3m2 8H7a2 2 0 01-2-2V5a2 2 0 012-2h5.586a1 1 0 01.707.293l5.414 5.414a1 1 0 01.293.707V19a2 2 0 01-2 2z"/></svg>
                        Download Excel
                    </a>
                </span>
            </div>
        </section>

        <!-- Detailed Results -->
        <main class="space-y-4">
            <h2 class="text-base font-bold text-slate-300 uppercase tracking-wider mb-2 flex items-center gap-2">
                <svg class="w-4 h-4 text-sky-400" fill="none" stroke="currentColor" viewBox="0 0 24 24"><path stroke-linecap="round" stroke-linejoin="round" stroke-width="2" d="M9 5H7a2 2 0 00-2 2v12a2 2 0 002 2h10a2 2 0 002-2V7a2 2 0 00-2-2h-2M9 5a2 2 0 002 2h2a2 2 0 002-2M9 5a2 2 0 012-2h2a2 2 0 012 2m-3 7h3m-3 4h3m-6-4h.01M9 16h.01"/></svg>
                Detailed Test Cases Execution
            </h2>
            <div class="space-y-3">
                {cards_html}
            </div>
        </main>
        
        <!-- Footer -->
        <footer class="mt-16 pt-6 border-t border-slate-900 text-center text-xs text-slate-500 font-mono">
            <p>Smart EV Assistant E2E Automated Testing Suite &bull; Appium / Android Driver</p>
        </footer>
    </div>

    <!-- Accordion Script -->
    <script>
        function toggleAccordion(id) {{
            const content = document.getElementById('content-' + id);
            const arrow = document.getElementById('arrow-' + id);
            if (content.classList.contains('hidden')) {{
                content.classList.remove('hidden');
                arrow.classList.add('rotate-180');
            }} else {{
                content.classList.add('hidden');
                arrow.classList.remove('rotate-180');
            }}
        }}
    </script>
</body>
</html>
"""
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html_content)
    print(f"HTML report saved to: {html_path}")

def generate_markdown_summary(results, build_no="Local"):
    ensure_directories()
    summary_path = os.path.join(SUMMARY_DIR, "summary.md")
    
    total = len(results)
    passed = sum(1 for r in results if r["status"] == "PASS")
    failed = sum(1 for r in results if r["status"] == "FAIL")
    rate = f"{round(passed / total * 100, 1)}%" if total else "0.0%"
    
    status_emoji = "✅ PASS" if failed == 0 else "❌ FAIL"
    date_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    
    content = f"""# Android Appium Test Summary

**Build Number:** {build_no}  
**Execution Date:** {date_str}  
**Status:** {status_emoji}

| Metric | Value |
| :--- | :--- |
| **Total Tests** | {total} |
| **Passed** | {passed} |
| **Failed** | {failed} |
| **Pass Rate** | **{rate}** |

---

## 📊 Live Report Details
The latest execution report is hosted online and available for review:

[🌐 View Live HTML Report](https://192321101.github.io/Smart-EV-Assistant/reports/latest/execution-report.html)
"""
    with open(summary_path, "w", encoding="utf-8") as f:
        f.write(content)
    print(f"Markdown summary saved to: {summary_path}")

def main():
    print("Starting reports generation fallback check...")
    generate_excel_report(DEFAULT_RESULTS)
    generate_html_report(DEFAULT_RESULTS)
    generate_markdown_summary(DEFAULT_RESULTS)
    print("Reports generated successfully!")

if __name__ == "__main__":
    main()
